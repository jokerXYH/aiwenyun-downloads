from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import analyze
import excel
import matrix
import net
import pipeline
import token_source


def sample_token(ts: int = 1792944444464) -> str:
    return '4041-0-454429-' + 'a' * 32 + '-%d-24.8.8_5.72.216-PC-ai' % ts


class TokenTests(unittest.TestCase):
    def test_strict_regex_and_metadata(self):
        token = sample_token()
        self.assertTrue(token_source.TOKEN_RE.fullmatch(token))
        meta = token_source.parse_token_metadata(token)
        self.assertEqual(meta['user_id'], '454429')
        self.assertEqual(meta['version'], '24.8.8_5.72.216')
        self.assertEqual(meta['platform'], 'ai')

    def test_extract_ascii(self):
        token = sample_token()
        data = b'xx\x00' + token.encode() + b'\x00yy'
        found = token_source._extract_from_bytes(data, 123, 0x1000)
        self.assertIn(token, [item.value for item in found])

    def test_extract_utf16(self):
        token = sample_token()
        data = ('\x00prefix\x00' + token + '\x00suffix\x00').encode('utf-16le')
        found = token_source._extract_from_bytes(data, 123, 0x2000)
        self.assertIn(token, [item.value for item in found])

    def test_no_false_positive(self):
        found = token_source._extract_from_bytes(b'abc-PC-def 123456', 1, 0)
        self.assertEqual(found, [])


class AnalyzeTests(unittest.TestCase):
    def test_invite_structured(self):
        self.assertEqual(analyze.invite_code_from_payload(
            {'code': 0, 'obj': {'code': '457288'}}), '457288')
        self.assertEqual(analyze.invite_code_from_payload(
            {'obj': [{'inviteCode': '１２３４５６'}]}), '123456')
        self.assertEqual(analyze.invite_code_from_payload(
            {'obj': {'text': '邀请码：654321'}}), '654321')

    def test_invite_does_not_use_plain_ids(self):
        self.assertEqual(analyze.invite_code_from_payload(
            {'code': 0, 'obj': {'id': 123456, 'ts': 1792944444464}}), '')

    def test_invite_lenient_shapes(self):
        """服务端换字段/包一层时也能从 obj 里挑出唯一的 6 位码。"""
        self.assertEqual(analyze.invite_code_from_payload(
            {'code': 0, 'obj': '457288'}), '457288')
        self.assertEqual(analyze.invite_code_from_payload(
            {'code': 0, 'obj': {'value': '457288'}}), '457288')
        self.assertEqual(analyze.invite_code_from_payload(
            {'code': 0, 'obj': {'id': 123456, 'value': '457288'}}), '457288')
        self.assertEqual(analyze.invite_code_from_payload(
            {'code': 0, 'obj': {'detail': {'code': '457288'}}}), '457288')
        self.assertEqual(analyze.invite_code_from_payload(
            {'code': 0, 'obj': {'a': '111111', 'b': '222222'}}), '')

    def test_timetable_rows(self):
        rows = analyze.timetable_rows({
            'obj': {'total': 2, 'list': [
                {'taskId': 1, 'topic': 'A-8', 'startTime': 100},
                {'id': 2, 'title': 'B', 'beginTime': 200},
            ]}})
        self.assertEqual([row['id'] for row in rows], [1, 2])
        self.assertEqual(rows[0]['topic'], 'A-8')
        self.assertEqual(rows[0]['begin'], 100)

    def test_live_info_groups(self):
        groups = analyze.normalize_groups({'1': 'Aβ1', '2': 'Aβ3'})
        self.assertEqual(groups, ['Aβ1', 'Aβ3'])


class MatrixTests(unittest.TestCase):
    def test_lesson_and_course(self):
        self.assertEqual(matrix.lesson_from_topic('26秋7MAA12（数形许）-8'), 8)
        self.assertEqual(matrix.lesson_from_topic('26秋7MAA12（数形许）-8β1'), 8)
        self.assertEqual(matrix.lesson_from_topic('第9节'), 9)
        self.assertEqual(matrix.course_from_topic('26秋7MAA12（数形许）-8β1'),
                         '26秋7MAA12（数形许）')

    def test_matrix_and_missing_lessons(self):
        items = matrix.normalize_items([
            {'topic': 'B-2', 'code': '222222', 'begin': 20},
            {'topic': 'B-1', 'code': '111111', 'begin': 10},
            {'topic': 'A', 'code': '333333', 'begin': 30},
            {'topic': 'A', 'code': '444444', 'begin': 40},
        ])
        result = matrix.build_matrices(items)[0]
        self.assertEqual(result.labels, ['A', 'B'])
        self.assertEqual(result.lessons, [1, 2])
        self.assertEqual(result.code('B', 1), '111111')
        self.assertEqual(result.code('B', 2), '222222')
        self.assertEqual(result.code('A', 1), '333333')
        self.assertEqual(result.code('A', 2), '444444')

    def test_conflict_keeps_first(self):
        items = matrix.normalize_items([
            {'topic': 'A-1', 'code': '111111', 'begin': 1},
            {'topic': 'A-1', 'code': '222222', 'begin': 2},
        ])
        result = matrix.build_matrices(items)[0]
        self.assertEqual(result.code('A', 1), '111111')
        self.assertEqual(len(result.conflicts), 1)


class NetTests(unittest.TestCase):
    def test_redact_token(self):
        token = sample_token()
        text = 'invalid access_token:%s 请重试' % token
        cleaned = net.redact(text)
        self.assertNotIn(token, cleaned)
        self.assertIn('***', cleaned)

    def test_token_bucket(self):
        bucket = net.TokenBucket(rate=100, burst=5)
        started = time.monotonic()
        for _ in range(15):
            bucket.acquire()
        self.assertLess(time.monotonic() - started, 1.0)


class ExcelTests(unittest.TestCase):
    def test_write_and_load(self):
        items = matrix.normalize_items([
            {'topic': '26秋7MAA12（数形许）-8', 'code': '012345', 'begin': 1},
            {'topic': '26秋7MAA12（数形许）-9', 'code': '654321', 'begin': 2},
        ])
        result = matrix.build_matrices(items)[0]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'out.xlsx'
            actual = excel.write_workbook(
                path, [result],
                details=[{'date': '2026-10-31 10:20',
                          'classId': 1, 'topic': 'x-8', 'course': 'x',
                          'lesson': 8, 'code': '012345', 'note': '',
                          'groups': ['G1']}],
                metadata={'client_version': 'v', 'base_url': 'https://x/'})
            self.assertTrue(actual.is_file())
            from openpyxl import load_workbook
            workbook = load_workbook(actual)
            self.assertEqual(workbook.sheetnames, ['邀请码', '明细'])
            sheet = workbook['邀请码']
            self.assertEqual(sheet['A1'].value, '班级')
            self.assertEqual(sheet['B1'].value, '第8节')
            self.assertEqual(sheet['B2'].value, '012345')
            self.assertEqual(sheet['B2'].number_format, '@')
            self.assertEqual(sheet.freeze_panes, 'B2')
            detail = workbook['明细']
            self.assertEqual(detail.max_row, 2)
            self.assertEqual(detail['F2'].value, '012345')


class FakeSession:
    def __init__(self, codes):
        self.codes = codes
        self.expired_once = False

    def fetch_invite_code(self, class_id, deadline=None):
        if self.expired_once:
            self.expired_once = False
            raise net.TokenExpired('expired')
        return self.codes.get(class_id, '')


class FakeState:
    def __init__(self, codes):
        self.session = FakeSession(codes)
        self.generation = 0
        self.refresh_count = 0

    def current(self):
        return self.session, self.generation

    def refresh(self, seen):
        self.refresh_count += 1
        self.generation += 1
        self.session.expired_once = False
        return self.session, self.generation


class PipelineTests(unittest.TestCase):
    def test_fetch_invites_preserves_order(self):
        rows = [{'id': i, 'topic': 'A-%d' % i, 'begin': i} for i in range(1, 6)]
        codes = {i: '%06d' % i for i in range(1, 6)}
        state = FakeState(codes)
        options = pipeline.ExportOptions(workers=3, rate=100, include_details=False)
        progress = []
        results, canceled = pipeline.fetch_invites(
            state, rows, options, on_progress=lambda d, t, r: progress.append(d))
        self.assertFalse(canceled)
        self.assertEqual([item['code'] for item in results],
                         ['000001', '000002', '000003', '000004', '000005'])
        self.assertEqual(progress[-1], 5)

    def test_token_refresh_retry(self):
        rows = [{'id': 1, 'topic': 'A-1', 'begin': 1}]
        codes = {1: '123456'}
        state = FakeState(codes)
        state.session.expired_once = True
        options = pipeline.ExportOptions(workers=1, rate=100)
        results, canceled = pipeline.fetch_invites(state, rows, options)
        self.assertFalse(canceled)
        self.assertEqual(results[0]['code'], '123456')
        self.assertEqual(state.refresh_count, 1)

    def test_cancel_before_start(self):
        rows = [{'id': 1, 'topic': 'A-1', 'begin': 1}]
        state = FakeState({1: '123456'})
        cancel = threading.Event()
        cancel.set()
        options = pipeline.ExportOptions(workers=1, rate=100)
        results, canceled = pipeline.fetch_invites(
            state, rows, options, cancel_event=cancel)
        self.assertTrue(canceled)
        self.assertEqual(results[0]['code'], '')
        self.assertIn('取消', results[0]['note'])

    def test_stall_watchdog_returns_partial(self):
        """一个请求永远不返回时，不能把整次导出拖死。"""
        rows = [{'id': 1, 'topic': 'A-1', 'begin': 1},
                {'id': 2, 'topic': 'A-2', 'begin': 2}]
        state = FakeState({1: '123456'})
        release = threading.Event()
        blocked = {'value': False}

        def fake_fetch(fake_state, row, bucket, cancel_event,
                       cooldown=None, class_deadline=None, max_attempts=3):
            if row['id'] == 2 and not blocked['value']:
                blocked['value'] = True
                release.wait(5)
            record = pipeline._fallback_record(row, '')
            if row['id'] == 1:
                record['code'] = '123456'
            return record

        options = pipeline.ExportOptions(
            workers=2, rate=100, stall_timeout=0.4, retry_budget=0)
        started = time.monotonic()
        results, _ = pipeline.fetch_invites(
            state, rows, options, fetch_one=fake_fetch)
        elapsed = time.monotonic() - started
        release.set()
        time.sleep(0.1)
        self.assertLess(elapsed, 3.0)
        self.assertEqual(results[0]['code'], '123456')
        self.assertEqual(results[1]['code'], '')
        self.assertIn('超时', results[1]['note'])

    def test_second_pass_retries_failures(self):
        rows = [{'id': 1, 'topic': 'A-1', 'begin': 1}]
        state = FakeState({1: '654321'})
        calls = {'count': 0}

        def fake_fetch(fake_state, row, bucket, cancel_event,
                       cooldown=None, class_deadline=None, max_attempts=3):
            calls['count'] += 1
            record = pipeline._fallback_record(row, '网络失败')
            if calls['count'] >= 2:
                record['code'] = '654321'
                record['note'] = ''
            return record

        stats = {}
        options = pipeline.ExportOptions(
            workers=1, rate=100, retry_budget=5)
        results, _ = pipeline.fetch_invites(
            state, rows, options, fetch_one=fake_fetch,
            on_stats=lambda value: stats.update(value))
        self.assertEqual(calls['count'], 2)
        self.assertEqual(results[0]['code'], '654321')
        self.assertEqual(stats['retry']['attempted'], 1)
        self.assertEqual(stats['retry']['recovered'], 1)

    def test_cooldown_pauses_all_workers(self):
        cooldown = pipeline.Cooldown()
        cooldown.pause(0.2)
        started = time.monotonic()
        cooldown.wait()
        self.assertGreaterEqual(time.monotonic() - started, 0.18)

    def test_no_client(self):
        original_running = pipeline.client_running
        original_platform = pipeline.platform_name
        pipeline.client_running = lambda: False
        pipeline.platform_name = lambda: 'win32'
        try:
            with self.assertRaises(pipeline.NoClientError):
                pipeline.run_export(pipeline.ExportOptions())
        finally:
            pipeline.client_running = original_running
            pipeline.platform_name = original_platform


class MacLaunchTests(unittest.TestCase):
    def test_launch_app_uses_open_with_debug_args(self):
        import token_macos

        app = Path('/Applications/AiCloud.app')
        with patch.object(
                token_macos, '_bundle_executable',
                return_value=app / 'Contents/MacOS/AiCloud'), \
                patch.object(token_macos.time, 'sleep'), \
                patch.object(token_macos.subprocess, 'run') as run:
            run.return_value = subprocess.CompletedProcess(
                args=[], returncode=0, stdout='', stderr='')
            token_macos.launch_app(app)

        command = run.call_args.args[0]
        self.assertEqual(command[:4], ['open', '-a', str(app), '--args'])
        self.assertIn('--remote-debugging-address=127.0.0.1', command)
        self.assertIn('--remote-debugging-port=9222', command)

    def test_restore_app_uses_open_without_debug_args(self):
        import token_macos

        app = Path('/Applications/AiCloud.app')
        with patch.object(token_macos, 'quit_app') as quit_app, \
                patch.object(token_macos.subprocess, 'run') as run:
            token_macos.restore_app(app)

        quit_app.assert_called_once_with(app)
        self.assertEqual(run.call_args.args[0], ['open', '-a', str(app)])

    def test_pipeline_restores_macos_client(self):
        import token_macos

        with patch.object(pipeline.sys, 'platform', 'darwin'), \
                patch.object(token_macos, 'restore_app') as restore:
            pipeline._restore_macos_client(
                {'restarted': True, 'app': '/Applications/AiCloud.app'}, None)

        restore.assert_called_once_with(Path('/Applications/AiCloud.app'))


if __name__ == '__main__':
    unittest.main(verbosity=2)
